"""
Документи: Word, PDF, TXT, Excel и PowerPoint от опростен markdown.

Текстът (написан от orion/writer.py или даден от сър) е в този вид:

    # Заглавие                 ## Раздел            ### Подраздел
    - точка                    1. номерирана точка
    **удебелено**  *наклонено*
    | колона | колона |        (таблица, първият ред — заглавия)
    >> ред вдясно              (адресатът на молба: „ДО Директора на…“)
    [нова страница]

Всеки ред е отделен абзац — така писмата и молбите запазват вида си.
Файловете отиват в Документи\\Орион (или в папката, която сър каже), а списъкът с
последните документи е в memory/documents.json — за „отвори го“, „добави към него“.
"""
import json
import os
import re
import threading
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from . import folders

BASE_DIR = Path(__file__).resolve().parent.parent
RECENT_FILE = BASE_DIR / "memory" / "documents.json"
FONTS = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"

# Как сър нарича формата -> разширение.
FORMATS = {
    "word": ".docx", "docx": ".docx", "doc": ".docx", "ворд": ".docx", "уърд": ".docx",
    "pdf": ".pdf", "пдф": ".pdf",
    "txt": ".txt", "текст": ".txt", "текстов": ".txt", "бележник": ".txt", "notepad": ".txt",
    "md": ".md", "markdown": ".md",
}


def extension(file_format: str, default: str = ".docx") -> str:
    key = (file_format or "").lower().strip(" .")
    return next((ext for name, ext in FORMATS.items() if name == key or name in key.split()), default)


# --- Разбор на текста -------------------------------------------------------------------------
@dataclass
class Block:
    kind: str                 # title, heading, para, right, bullet, number, table, pagebreak
    text: str = ""
    level: int = 0            # за heading (2, 3) и за вложени точки (0, 1)
    number: str = ""          # „3.“ за номерираните точки
    rows: list[list[str]] = field(default_factory=list)


_TABLE_SEPARATOR = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def parse(markdown: str) -> list[Block]:
    blocks: list[Block] = []
    for raw in markdown.replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped or stripped.startswith("```") or re.fullmatch(r"[-*_]{3,}", stripped):
            continue
        if stripped.lower() in ("[нова страница]", "\\pagebreak", "<pagebreak>"):
            blocks.append(Block("pagebreak"))
            continue
        if stripped.startswith("|") and stripped.count("|") >= 2:
            if _TABLE_SEPARATOR.match(stripped):
                continue
            if blocks and blocks[-1].kind == "table":
                blocks[-1].rows.append(_cells(stripped))
            else:
                blocks.append(Block("table", rows=[_cells(stripped)]))
            continue
        match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if match:
            level = len(match.group(1))
            text = match.group(2).strip().strip("#").strip()
            has_title = any(b.kind == "title" for b in blocks)
            blocks.append(Block("title", text) if level == 1 and not has_title
                          else Block("heading", text, level=min(max(level, 2), 3)))
            continue
        if stripped.startswith(">>"):
            blocks.append(Block("right", stripped[2:].strip()))
            continue
        indent = len(line) - len(line.lstrip())
        match = re.match(r"^[-*•]\s+(.*)$", stripped)
        if match:
            blocks.append(Block("bullet", match.group(1), level=1 if indent >= 2 else 0))
            continue
        match = re.match(r"^(\d{1,3})[.)]\s+(.*)$", stripped)
        if match:
            blocks.append(Block("number", match.group(2), level=1 if indent >= 2 else 0,
                                number=f"{match.group(1)}."))
            continue
        if stripped.startswith(">"):  # обикновен цитат
            stripped = stripped.lstrip("> ").strip()
        blocks.append(Block("para", stripped))
    # Табличните редове с по-малко клетки се допълват — иначе Word и PDF не ги подреждат.
    for block in blocks:
        if block.kind == "table":
            width = max(len(r) for r in block.rows)
            block.rows = [r + [""] * (width - len(r)) for r in block.rows]
    return blocks


def plain(text: str) -> str:
    """Текстът без ** и * (за TXT, заглавия и имена на файлове)."""
    return re.sub(r"\*\*(.+?)\*\*|(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", lambda m: m.group(1) or m.group(2), text)


def _runs(text: str) -> list[tuple[str, bool, bool]]:
    """„**а** б *в*“ -> [(„а“, удебелено, наклонено), …]."""
    parts = re.split(r"(\*\*.+?\*\*|(?<!\*)\*(?!\s).+?(?<!\s)\*(?!\*))", text)
    runs = []
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            runs.append((part[2:-2], True, False))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            runs.append((part[1:-1], False, True))
        else:
            runs.append((part, False, False))
    return runs


def title_of(blocks: list[Block], fallback: str = "") -> str:
    for block in blocks:
        if block.kind in ("title", "heading") and block.text.strip():
            return plain(block.text).strip()
    for block in blocks:
        if block.kind in ("para", "bullet", "number") and block.text.strip():
            words = plain(block.text).split()
            return " ".join(words[:8])
    return fallback


def word_count(markdown: str) -> int:
    return len(re.findall(r"\w+", plain(markdown)))


# --- Къде се записва ------------------------------------------------------------------------
def default_folder() -> Path:
    return folders.known().get("документи", Path.home() / "Documents") / "Орион"


def target_path(title: str, ext: str, location: str = "") -> Path:
    """Свободно име на файл: „Молба за отпуск.docx“, а ако вече има — „Молба за отпуск (2).docx“."""
    folder = folders.resolve(location) if location and location.strip() else None
    if location and location.strip() and not folder:
        raise FileNotFoundError(f"не намирам папка „{location}“")
    folder = folder or default_folder()
    folder.mkdir(parents=True, exist_ok=True)
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", plain(title)).strip(" .")[:80].strip()
    name = name or f"Документ {datetime.now():%d.%m.%Y %H-%M}"
    path = folder / f"{name}{ext}"
    number = 2
    while path.exists():
        path = folder / f"{name} ({number}){ext}"
        number += 1
    return path


_recent_lock = threading.Lock()


def recent() -> list[Path]:
    """Документите, които Орион е създал — последният е първи. Само съществуващите."""
    try:
        items = json.loads(RECENT_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return []
    return [Path(p) for p in items if Path(p).exists()]


def remember(path: Path) -> None:
    with _recent_lock:
        items = [str(path)] + [str(p) for p in recent() if p != path]
        RECENT_FILE.parent.mkdir(exist_ok=True)
        RECENT_FILE.write_text(json.dumps(items[:30], ensure_ascii=False, indent=1), encoding="utf-8")


def last() -> Path:
    items = recent()
    if not items:
        raise LookupError("още не съм създавал документи")
    return items[0]


# --- Word -------------------------------------------------------------------------------------
def _docx_runs(paragraph, text: str, size=None, bold: bool = False) -> None:
    from docx.shared import Pt
    for chunk, is_bold, is_italic in _runs(text):
        run = paragraph.add_run(chunk)
        run.bold = bold or is_bold or None
        run.italic = is_italic or None
        if size:
            run.font.size = Pt(size)


def _docx_blocks(doc, blocks: list[Block]) -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    accent = RGBColor(0x1F, 0x3A, 0x5F)
    for block in blocks:
        if block.kind == "title":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(14)
            _docx_runs(p, block.text, size=16, bold=True)
            for run in p.runs:
                run.font.color.rgb = accent
        elif block.kind == "heading":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            _docx_runs(p, block.text, size=13.5 if block.level == 2 else 12, bold=True)
            for run in p.runs:
                run.font.color.rgb = accent
        elif block.kind == "right":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_after = Pt(0)
            _docx_runs(p, block.text)
        elif block.kind == "bullet":
            p = doc.add_paragraph(style="List Bullet 2" if block.level else "List Bullet")
            p.paragraph_format.space_after = Pt(2)
            _docx_runs(p, block.text)
        elif block.kind == "number":
            # Номерът е текст: вградената номерация на Word продължава от предишния списък.
            # \u041e\u0442\u0441\u0442\u044a\u043f\u044a\u0442 \u0435 \u043a\u0430\u0442\u043e \u043d\u0430 \u0442\u043e\u0447\u043a\u0438\u0442\u0435 (\u201eList Bullet\u201c), \u0430 \u0442\u0430\u0431\u0443\u043b\u0430\u0446\u0438\u044f\u0442\u0430 \u043f\u043e\u0434\u0440\u0430\u0432\u043d\u044f\u0432\u0430 \u0442\u0435\u043a\u0441\u0442\u0430.
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.63 * (block.level + 1))
            p.paragraph_format.first_line_indent = Cm(-0.63)
            p.paragraph_format.space_after = Pt(2)
            _docx_runs(p, f"{block.number}\t{block.text}")
        elif block.kind == "table":
            table = doc.add_table(rows=len(block.rows), cols=len(block.rows[0]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, row in enumerate(block.rows):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = ""
                    _docx_runs(cell.paragraphs[0], value, bold=(r == 0))
            doc.add_paragraph()
        elif block.kind == "pagebreak":
            doc.add_page_break()
        else:
            p = doc.add_paragraph()
            _docx_runs(p, block.text)


def to_docx(blocks: list[Block], path: Path, title: str = "") -> None:
    from docx import Document
    from docx.shared import Cm, Pt

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)  # A4 (по подразбиране е US Letter)
    section.left_margin = section.right_margin = Cm(2.5)
    section.top_margin = section.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    _docx_blocks(doc, blocks)
    doc.core_properties.title = title
    doc.core_properties.author = "Орион"
    doc.save(path)


# --- PDF --------------------------------------------------------------------------------------
def _pdf_text(text: str) -> str:
    """Нашият markdown -> този на fpdf2 (**удебелено**, __наклонено__); „--“ там е подчертаване."""
    out = []
    for chunk, bold, italic in _runs(text):
        chunk = chunk.replace("--", "–").replace("__", "_").replace("**", "")
        out.append(f"**{chunk}**" if bold else f"__{chunk}__" if italic else chunk)
    return "".join(out)


def to_pdf(blocks: list[Block], path: Path, title: str = "") -> None:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF(format="A4")
    pdf.set_margins(22, 20, 22)
    pdf.set_auto_page_break(True, margin=20)
    pdf.set_title(title)
    pdf.set_author("Орион")
    for style, file in (("", "arial.ttf"), ("B", "arialbd.ttf"), ("I", "ariali.ttf"), ("BI", "arialbi.ttf")):
        pdf.add_font("Main", style, str(FONTS / file))
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin
    line = dict(new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    for block in blocks:
        pdf.set_text_color(0, 0, 0)
        if block.kind == "title":
            pdf.set_font("Main", "B", 16)
            pdf.set_text_color(31, 58, 95)
            pdf.ln(2)
            pdf.multi_cell(width, 8, plain(block.text), align="C", **line)
            pdf.ln(5)
        elif block.kind == "heading":
            pdf.set_font("Main", "B", 13 if block.level == 2 else 12)
            pdf.set_text_color(31, 58, 95)
            pdf.ln(3)
            pdf.multi_cell(width, 7, plain(block.text), **line)
            pdf.ln(1)
        elif block.kind == "right":
            pdf.set_font("Main", "", 11.5)
            pdf.multi_cell(width, 6, _pdf_text(block.text), align="R", markdown=True, **line)
        elif block.kind in ("bullet", "number"):
            pdf.set_font("Main", "", 11.5)
            indent = 5 + 6 * block.level
            pdf.set_x(pdf.l_margin + indent)
            mark = "•" if block.kind == "bullet" else block.number
            pdf.cell(7, 6, mark)
            pdf.multi_cell(width - indent - 7, 6, _pdf_text(block.text), markdown=True, **line)
            pdf.ln(0.8)
        elif block.kind == "table":
            pdf.set_font("Main", "", 10.5)
            pdf.ln(1)
            with pdf.table(first_row_as_headings=True, line_height=6.5, padding=1.5) as table:
                for row in block.rows:
                    cells = table.row()
                    for value in row:
                        cells.cell(plain(value))
            pdf.ln(3)
        elif block.kind == "pagebreak":
            pdf.add_page()
        else:
            pdf.set_font("Main", "", 11.5)
            pdf.multi_cell(width, 6, _pdf_text(block.text), markdown=True, **line)
            pdf.ln(2)
    pdf.output(str(path))


# --- TXT и Markdown ------------------------------------------------------------------------------
def to_text(blocks: list[Block]) -> str:
    lines: list[str] = []
    for block in blocks:
        if block.kind == "title":
            lines += [plain(block.text).upper(), ""]
        elif block.kind == "heading":
            lines += ["", plain(block.text), ""]
        elif block.kind == "right":
            lines.append(plain(block.text).rjust(70))
        elif block.kind == "bullet":
            lines.append("    " * block.level + "• " + plain(block.text))
        elif block.kind == "number":
            lines.append("    " * block.level + f"{block.number} " + plain(block.text))
        elif block.kind == "table":
            widths = [max(len(plain(r[c])) for r in block.rows) for c in range(len(block.rows[0]))]
            for r, row in enumerate(block.rows):
                lines.append(" | ".join(plain(v).ljust(w) for v, w in zip(row, widths)))
                if r == 0:
                    lines.append("-+-".join("-" * w for w in widths))
            lines.append("")
        elif block.kind == "pagebreak":
            lines += ["", "\f"]
        else:
            lines.append(plain(block.text))
    return "\n".join(lines).strip("\n") + "\n"


def save(markdown: str, path: Path, title: str = "") -> None:
    """Записва текста във формата на разширението на `path`."""
    blocks = parse(markdown)
    if not blocks:
        raise ValueError("документът е празен")
    ext = path.suffix.lower()
    if ext == ".docx":
        to_docx(blocks, path, title)
    elif ext == ".pdf":
        to_pdf(blocks, path, title)
    elif ext == ".md":
        path.write_text(markdown.strip() + "\n", encoding="utf-8")
    else:
        path.write_text(to_text(blocks), encoding="utf-8-sig")  # с BOM — Notepad чете кирилицата
    remember(path)


def append(path: Path, markdown: str) -> None:
    """Добавя текст в края на съществуващ Word, TXT или Markdown файл."""
    ext = path.suffix.lower()
    blocks = parse(markdown)
    if not blocks:
        raise ValueError("няма какво да добавя")
    if ext == ".docx":
        from docx import Document
        doc = Document(path)
        _docx_blocks(doc, blocks)
        doc.save(path)
    elif ext == ".md":
        with open(path, "a", encoding="utf-8") as f:
            f.write("\n" + markdown.strip() + "\n")
    elif ext == ".txt":
        with open(path, "a", encoding="utf-8") as f:
            f.write("\n" + to_text(blocks))
    else:
        raise ValueError(f"не мога да добавям в {ext} файл — мога във Word, TXT и Markdown")
    remember(path)


# --- Excel ------------------------------------------------------------------------------------
_EXCEL_NAMES = {"СУМА": "SUM", "СРЕДНО": "AVERAGE", "МАКС": "MAX", "МИН": "MIN", "БРОЙ": "COUNT",
                "АКО": "IF", "ЗАКРЪГЛИ": "ROUND", "ДНЕС": "TODAY"}
_NUMBER_RE = re.compile(r"^(?P<num>-?\d{1,3}(?:[ \u00a0\u202f]\d{3})+(?:[.,]\d+)?|-?\d+(?:[.,]\d+)?)"
                        r"\s*(?P<unit>%|лв\.?|лева|€|евро|eur|\$|usd)?$", re.IGNORECASE)
_DATE_RE = re.compile(r"^(\d{1,2})\.(\d{1,2})\.(\d{4})(?:\s*г\.?)?$")


def _excel_value(text: str):
    """Текст от таблицата -> (стойност, числов формат или None)."""
    value = plain(text).strip()
    if value.startswith("="):
        formula = value.replace(";", ",")
        for bg, en in _EXCEL_NAMES.items():
            formula = re.sub(rf"\b{bg}\(", f"{en}(", formula, flags=re.IGNORECASE)
        return formula, None
    match = _NUMBER_RE.match(value)
    if match:
        raw = re.sub(r"[ \u00a0\u202f]", "", match["num"])
        if "," in raw and "." not in raw:
            raw = raw.replace(",", ".")
        elif "," in raw:
            raw = raw.replace(",", "")
        number = float(raw)
        number = int(number) if number.is_integer() and "." not in raw else number
        unit = (match["unit"] or "").lower()
        if unit == "%":
            return number / 100, "0%" if float(number).is_integer() else "0.0%"
        if unit.startswith("лв") or unit == "лева":
            return number, '#,##0.00 "лв."'
        if unit in ("€", "евро", "eur"):
            return number, '#,##0.00 "€"'
        if unit in ("$", "usd"):
            return number, '"$"#,##0.00'
        return number, "#,##0.00" if isinstance(number, float) else None
    match = _DATE_RE.match(value)
    if match:
        try:
            return datetime(int(match[3]), int(match[2]), int(match[1])), "DD.MM.YYYY"
        except ValueError:
            pass
    return value, None


def _column_format(ws, column: int) -> None:
    """Еднакъв вид на числата в колоната: „Сума (лв.)“ — с „лв.“, ако има дробни — с две цифри
    след запетаята, процентите — с %. И сумата най-долу (формулата) е в същия вид."""
    from collections import Counter
    header = str(ws.cell(row=1, column=column).value or "").lower()
    cells = [ws.cell(row=r, column=column) for r in range(2, ws.max_row + 1)]
    numbers = [c for c in cells if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]
    formulas = [c for c in cells if isinstance(c.value, str) and c.value.startswith("=")]
    if not numbers and not formulas:
        return
    formats = Counter(c.number_format for c in numbers if c.number_format != "General")
    if "лв" in header:
        chosen = '#,##0.00 "лв."'
    elif "€" in header or "евро" in header or "eur" in header:
        chosen = '#,##0.00 "€"'
    elif "$" in header or "долар" in header or "usd" in header:
        chosen = '"$"#,##0.00'
    elif formats:
        chosen = formats.most_common(1)[0][0]
    else:
        return
    if chosen == "0%" and any(round(c.value * 100, 6) % 1 for c in numbers):
        chosen = "0.0%"
    for cell in numbers + formulas:
        cell.number_format = chosen


def sheets_from(markdown: str) -> list[tuple[str, list[list[str]]]]:
    """[(име на лист, редове)] — всяко „## Име“ започва нов лист; редовете са от таблиците
    с „|“ (или редове с „;“ / табулация, ако таблица няма)."""
    sheets: list[tuple[str, list[list[str]]]] = []
    name = ""
    for block in parse(markdown):
        if block.kind == "heading":
            name = plain(block.text)
        elif block.kind == "table":
            sheets.append((name or f"Лист{len(sheets) + 1}", block.rows))
            name = ""
    if not sheets:  # CSV: „Име;Сума“
        rows = [re.split(r"\s*[;\t]\s*", line.strip()) for line in markdown.splitlines()
                if (";" in line or "\t" in line) and line.strip()]
        if rows:
            width = max(len(r) for r in rows)
            sheets.append(("Лист1", [r + [""] * (width - len(r)) for r in rows]))
    return sheets


def to_xlsx(sheets: list[tuple[str, list[list[str]]]], path: Path, title: str = "") -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    if not sheets:
        raise ValueError("няма таблица")
    workbook = Workbook()
    workbook.remove(workbook.active)
    thin = Side(style="thin", color="C9D3DD")
    used_names: set[str] = set()
    for sheet_name, rows in sheets:
        name = re.sub(r"[\[\]:*?/\\]", "", sheet_name)[:31].strip() or "Лист"
        while name in used_names:
            name = name[:28] + f" {len(used_names) + 1}"
        used_names.add(name)
        ws = workbook.create_sheet(name)
        for r, row in enumerate(rows, 1):
            for c, text in enumerate(row, 1):
                cell = ws.cell(row=r, column=c)
                if r == 1:
                    cell.value = plain(text)
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="1F3A5F")
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                else:
                    value, number_format = _excel_value(text)
                    cell.value = value
                    if number_format:
                        cell.number_format = number_format
                    if plain(str(row[0])).strip().lower().startswith(("общо", "всичко", "сума", "total")):
                        cell.font = Font(bold=True)  # редът с общата сума
                cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        for c in range(1, max(len(r) for r in rows) + 1):
            longest = max(len(plain(str(row[c - 1]))) if c - 1 < len(row) else 0 for row in rows)
            ws.column_dimensions[get_column_letter(c)].width = min(max(10, longest + 3), 60)
            _column_format(ws, c)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    workbook.properties.title = title
    workbook.properties.creator = "Орион"
    workbook.save(path)
    remember(path)


# --- PowerPoint -------------------------------------------------------------------------------
@dataclass
class Slide:
    title: str
    points: list[tuple[str, int]] = field(default_factory=list)  # (текст, ниво)
    notes: str = ""


def slides_from(markdown: str) -> tuple[str, str, list[Slide]]:
    """(заглавие, подзаглавие, слайдове) от „# Заглавие“, ред подзаглавие, „## Слайд“, „- точка“."""
    title, subtitle, slides = "", "", []
    for block in parse(markdown):
        text = plain(block.text)
        if block.kind == "title":
            title = text
        elif block.kind == "heading":
            slides.append(Slide(text))
        elif block.kind in ("bullet", "number") and slides:
            slides[-1].points.append((text, block.level))
        elif block.kind in ("para", "right"):
            if re.match(r"^(бележки|notes)\s*:", text, re.IGNORECASE) and slides:
                slides[-1].notes = text.split(":", 1)[1].strip()
            elif not slides and not subtitle:
                subtitle = text
            elif slides:
                slides[-1].points.append((text, 0))
        elif block.kind == "table" and slides:
            slides[-1].points += [(" · ".join(plain(c) for c in row if c), 0) for row in block.rows]
    return title, subtitle, slides


def to_pptx(title: str, subtitle: str, slides: list[Slide], path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    navy, accent, ink = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0x2C, 0x9C, 0xCB), RGBColor(0x22, 0x2B, 0x33)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9
    W, H = prs.slide_width, prs.slide_height

    def bar(slide, top, height, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, top, W, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()

    def style(frame, size, color, bold=False):
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                run.font.size, run.font.bold, run.font.name = Pt(size), bold, "Calibri"
                run.font.color.rgb = color

    cover = prs.slides.add_slide(prs.slide_layouts[0])
    bar(cover, 0, H, navy)
    bar(cover, Inches(4.9), Inches(0.08), accent)
    head, sub = cover.placeholders[0], cover.placeholders[1]
    for shape in (head, sub):
        shape.left, shape.width = Inches(0.9), W - Inches(1.8)
    head.top, head.height = Inches(2.2), Inches(2.4)
    sub.top, sub.height = Inches(5.2), Inches(1.2)
    head.text = title or "Презентация"
    sub.text = subtitle or datetime.now().strftime("%d.%m.%Y")
    style(head.text_frame, 44, RGBColor(0xFF, 0xFF, 0xFF), bold=True)
    style(sub.text_frame, 22, RGBColor(0xBC, 0xD9, 0xE8))
    # Правоъгълниците са добавени след полетата — местим ги отзад, за да не ги закриват.
    tree = cover.shapes._spTree
    for shape in [s for s in cover.shapes if not s.is_placeholder]:
        tree.remove(shape._element)
        tree.insert(2, shape._element)

    for item in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        bar(slide, 0, Inches(0.18), accent)
        head, body = slide.placeholders[0], slide.placeholders[1]
        head.left, head.top, head.width, head.height = Inches(0.8), Inches(0.45), W - Inches(1.6), Inches(1.2)
        body.left, body.top, body.width, body.height = Inches(0.8), Inches(1.8), W - Inches(1.6), H - Inches(2.4)
        head.text = item.title
        style(head.text_frame, 34, navy, bold=True)
        frame = body.text_frame
        frame.clear()
        points = item.points or [("", 0)]
        size = 26 if len(points) <= 4 else 22 if len(points) <= 6 else 18
        for i, (text, level) in enumerate(points):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = text
            paragraph.level = level
            paragraph.space_after = Pt(10)
        style(frame, size, ink)
        for paragraph in frame.paragraphs:
            if paragraph.level:
                for run in paragraph.runs:
                    run.font.size = Pt(size - 4)
        if item.notes:
            slide.notes_slide.notes_text_frame.text = item.notes
    prs.core_properties.title = title
    prs.core_properties.author = "Орион"
    prs.save(path)
    remember(path)


# --- PDF от Word, Excel и PowerPoint (през самите програми на Office) ---------------------------
def office_to_pdf(path: Path, timeout: float = 120) -> Path:
    """Превръща .docx/.xlsx/.pptx в PDF до оригинала. TXT и Markdown — без Office."""
    ext = path.suffix.lower()
    out = target_path(path.stem, ".pdf", str(path.parent))
    if ext in (".txt", ".md"):
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        to_pdf(parse(text), out, path.stem)
        remember(out)
        return out
    if ext not in (".docx", ".doc", ".rtf", ".odt", ".xlsx", ".xls", ".pptx", ".ppt"):
        raise ValueError(f"не мога да превърна {ext} в PDF")
    errors: list[Exception] = []

    def convert():
        import pythoncom
        import win32com.client
        pythoncom.CoInitialize()
        app = None
        try:
            if ext in (".docx", ".doc", ".rtf", ".odt"):
                app = win32com.client.DispatchEx("Word.Application")
                app.Visible, app.DisplayAlerts = False, 0
                document = app.Documents.Open(str(path), ReadOnly=True, AddToRecentFiles=False)
                document.ExportAsFixedFormat(str(out), 17)  # wdExportFormatPDF
                document.Close(False)
            elif ext in (".xlsx", ".xls"):
                app = win32com.client.DispatchEx("Excel.Application")
                app.Visible, app.DisplayAlerts = False, False
                workbook = app.Workbooks.Open(str(path), ReadOnly=True)
                workbook.ExportAsFixedFormat(0, str(out))  # xlTypePDF
                workbook.Close(False)
            else:
                app = win32com.client.DispatchEx("PowerPoint.Application")
                presentation = app.Presentations.Open(str(path), ReadOnly=True, WithWindow=False)
                presentation.SaveAs(str(out), 32)  # ppSaveAsPDF
                presentation.Close()
        except Exception as e:  # noqa: BLE001 — връща се на нишката, която чака
            errors.append(e)
        finally:
            if app is not None:
                try:
                    app.Quit()
                except Exception:  # noqa: BLE001
                    pass
            pythoncom.CoUninitialize()

    worker = threading.Thread(target=convert, daemon=True, name="pdf")
    worker.start()
    worker.join(timeout)
    if worker.is_alive():
        raise TimeoutError("Office не отговаря — може би чака отговор в свой прозорец")
    if errors:
        raise RuntimeError(f"Office не успя: {errors[0]}")
    if not out.exists():
        raise RuntimeError("Office не създаде PDF файла")
    remember(out)
    return out
